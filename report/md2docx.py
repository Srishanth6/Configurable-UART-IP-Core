#!/usr/bin/env python3
"""Tiny converter for my report markdown -> docx (headings, paragraphs, lists, tables, images, code)."""
import re, sys
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
src, out = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles["Normal"]; st.font.name = "DejaVu Serif"; st.font.size = Pt(10.5)
for s in doc.sections: s.left_margin = s.right_margin = Cm(2.2)

def runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"): p.add_run(part[2:-2]).bold = True
        elif part.startswith("*"): p.add_run(part[1:-1]).italic = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(9)
        else: p.add_run(part)

lines = open(src).read().splitlines(); i = 0
while i < len(lines):
    l = lines[i]
    if l.startswith("```"):
        i += 1; code = []
        while not lines[i].startswith("```"): code.append(lines[i]); i += 1
        p = doc.add_paragraph(); r = p.add_run("\n".join(code)); r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(8.5)
    elif l.startswith("#"):
        n = len(l) - len(l.lstrip("#")); h = doc.add_heading(l[n:].strip(), level=min(n - 1, 3) if n > 1 else 0)
    elif l.startswith("!["):
        m = re.match(r"!\[.*\]\((.+)\)", l); doc.add_picture(m.group(1), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif l.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[-| ]+\|$", lines[i]): rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        i -= 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, txt in enumerate(row):
                cell = t.cell(r, c); cell.text = ""; p = cell.paragraphs[0]; runs(p, txt)
                for rr in p.runs:
                    rr.font.size = Pt(8.5)
                    if r == 0: rr.bold = True
        doc.add_paragraph()
    elif l.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); runs(p, l[2:])
    elif re.match(r"^\d+\. ", l):
        p = doc.add_paragraph(style="List Number"); runs(p, l.split(" ", 1)[1])
    elif l.strip():
        p = doc.add_paragraph(); runs(p, l.rstrip(" "))
        if l.startswith("*Figure"): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    i += 1
doc.save(out)
